from io import BytesIO
from pathlib import Path
import fitz
import pandas as pd
from pptx import Presentation

class DocumentProcessingError(Exception):
    pass

def extract_text(filename: str, file_bytes: bytes) -> str:
    extension = Path(filename).suffix.lower()
    try:
        if extension == ".pdf":
            parts = []
            with fitz.open(stream=file_bytes, filetype="pdf") as doc:
                for i, page in enumerate(doc, start=1):
                    text = page.get_text("text").strip()
                    if text:
                        parts.append(f"[Page {i}]\n{text}")
            return "\n\n".join(parts)
        if extension == ".xlsx":
            book = pd.ExcelFile(BytesIO(file_bytes))
            parts = []
            for sheet in book.sheet_names:
                df = book.parse(sheet).dropna(how="all").dropna(axis=1, how="all")
                if not df.empty:
                    parts.append(f"[Worksheet: {sheet}]\n{df.to_csv(index=False)}")
            return "\n\n".join(parts)
        if extension == ".pptx":
            deck = Presentation(BytesIO(file_bytes))
            parts = []
            for i, slide in enumerate(deck.slides, start=1):
                texts = [s.text.strip() for s in slide.shapes if hasattr(s, "text") and s.text.strip()]
                if texts:
                    parts.append(f"[Slide {i}]\n" + "\n".join(texts))
            return "\n\n".join(parts)
    except Exception as exc:
        raise DocumentProcessingError(f"Could not process {filename}: {exc}") from exc
    raise DocumentProcessingError(f"Unsupported file type: {extension}")

